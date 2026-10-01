"""Build the interactive previews the file viewer shows.

Every sample in downloads/samples/ gets a preview at downloads/samples/previews/<key>.html:
real text and real cells, not a picture, so it stays sharp at any zoom.

  .xlsx  ->  an Excel-style grid per sheet: column letters, row numbers, the file's own
             column widths, row heights, fills, fonts, borders, merges, links and frozen panes.
  .docx  ->  Word's own filtered-HTML export (needs Windows with Word installed), set on a page.
  .html  ->  the issue itself, its own styles kept.

Run after any sample changes:   python tools/build_previews.py
Requires: openpyxl, pywin32 (for .docx).
"""
import datetime as dt
import html
import pathlib
import re
import shutil
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "downloads" / "samples"
OUT = SAMPLES / "previews"

FILES = {
    "chapter-settings": "Chapter_Settings_SAMPLE.docx",
    "chapter-plan": "Chapter_Plan_SAMPLE.docx",
    "registry": "CTC_Source_Registry_SAMPLE.xlsx",
    "events": "CTC_Events_SAMPLE.xlsx",
    "opportunities": "Opportunities_SAMPLE.xlsx",
    "city-resources": "CTC_City_Resources_SAMPLE.xlsx",
    "partner-map": "CTC_Partner_Map_SAMPLE.xlsx",
    "issue": "Portland_Climate_Tech_Aug_11-Aug_18.docx",
    "issue-html": "Portland_Climate_Tech_Aug_11-Aug_18.html",
}

# ---------------------------------------------------------------- xlsx
XL_CSS = """
:host { all: initial; }
.xl { font-family: Calibri, Carlito, "Segoe UI", Arial, sans-serif; font-size: 11pt; color: #000; background: #fff; display: inline-block; min-width: 100%; }
.sheet { display: none; }
.sheet.on { display: block; }
table { border-collapse: separate; border-spacing: 0; table-layout: fixed; background: #fff; }
td, th { padding: 2px 5px; overflow: hidden; white-space: nowrap; text-overflow: clip; vertical-align: bottom; box-sizing: border-box;
         border-right: 1px solid #e1e1e1; border-bottom: 1px solid #e1e1e1; line-height: 1.25; }
td.wrap { white-space: normal; overflow-wrap: anywhere; }
.nogrid td { border-color: transparent; }
th { background: #f3f3f3; color: #555; font: 400 10pt "Segoe UI", Arial, sans-serif; text-align: center; vertical-align: middle;
     border-right: 1px solid #d4d4d4; border-bottom: 1px solid #d4d4d4; position: sticky; z-index: 3; }
thead th { top: 0; height: 22px; }
tbody th { left: 0; z-index: 2; }
thead th.corner { left: 0; z-index: 4; }
td.frz { position: sticky; z-index: 1; }
a { color: #0563c1; text-decoration: underline; }
"""


def px_w(width):
    return int(round((width if width else 8.43) * 7 + 5))


def px_h(height):
    return int(round((height if height else 15) * 4 / 3))


def colour(c):
    if c is None or getattr(c, "type", None) != "rgb" or not isinstance(c.rgb, str):
        return None
    rgb = c.rgb[-6:]
    if c.rgb in ("00000000",) or not re.fullmatch(r"[0-9A-Fa-f]{6}", rgb):
        return None
    return "#" + rgb.lower()


def show(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, float):
        if "%" in (cell.number_format or ""):
            return f"{v * 100:.0f}%"
        return f"{v:g}"
    return str(v)


def cell_style(cell):
    s = []
    f = cell.fill
    if f is not None and f.fill_type == "solid":
        bg = colour(f.fgColor) or colour(f.start_color)
        if bg:
            s.append(f"background:{bg}")
    ft = cell.font
    if ft is not None:
        if ft.b:
            s.append("font-weight:700")
        if ft.i:
            s.append("font-style:italic")
        if ft.sz and float(ft.sz) != 11:
            s.append(f"font-size:{float(ft.sz):g}pt")
        col = colour(ft.color)
        if col:
            s.append(f"color:{col}")
        if ft.u:
            s.append("text-decoration:underline")
    al = cell.alignment
    if al is not None:
        if al.horizontal in ("center", "right", "left", "justify"):
            s.append(f"text-align:{al.horizontal}")
        elif al.horizontal == "centerContinuous":
            s.append("text-align:center")
        elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            s.append("text-align:right")
        if al.vertical in ("top", "center"):
            s.append("vertical-align:" + ("middle" if al.vertical == "center" else "top"))
    b = cell.border
    if b is not None:
        for side, name in ((b.left, "left"), (b.right, "right"), (b.top, "top"), (b.bottom, "bottom")):
            if side is not None and side.style:
                w = 2 if side.style in ("medium", "thick", "double") else 1
                s.append(f"border-{name}:{w}px solid {colour(side.color) or '#000'}")
    return ";".join(s)


def sheet_html(ws):
    max_r, max_c = ws.max_row, ws.max_column
    # trim trailing empty rows and columns
    while max_r > 1 and all(ws.cell(max_r, c).value is None for c in range(1, max_c + 1)):
        max_r -= 1
    while max_c > 1 and all(ws.cell(r, max_c).value is None for r in range(1, max_r + 1)):
        max_c -= 1
    # Excel stores widths per range of columns (min..max), keyed by the first letter only
    width_of = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for c in range(dim.min, dim.max + 1):
                width_of[c] = dim.width
    default_w = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.43
    widths = [px_w(width_of.get(c, default_w)) for c in range(1, max_c + 1)]
    heights = [px_h(ws.row_dimensions[r].height) for r in range(1, max_r + 1)]
    merged, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))
    frz_r = frz_c = 0
    if ws.freeze_panes:
        m = re.match(r"([A-Z]+)(\d+)", str(ws.freeze_panes))
        if m:
            frz_c = openpyxl.utils.column_index_from_string(m.group(1)) - 1
            frz_r = int(m.group(2)) - 1
    grid = "" if ws.sheet_view.showGridLines is not False else " nogrid"
    out = [f'<table class="{grid.strip()}" style="width:{44 + sum(widths)}px"><colgroup><col style="width:44px">']
    out += [f'<col style="width:{w}px">' for w in widths]
    out.append('</colgroup><thead><tr><th class="corner"></th>')
    left = 44
    lefts = []
    for c in range(1, max_c + 1):
        lefts.append(left)
        left += widths[c - 1]
        out.append(f"<th>{get_column_letter(c)}</th>")
    out.append("</tr></thead><tbody>")
    top = 22
    for r in range(1, max_r + 1):
        h = heights[r - 1]
        out.append(f'<tr style="height:{h}px"><th style="height:{h}px{";top:" + str(top) + "px;z-index:4" if r <= frz_r else ""}">{r}</th>')
        for c in range(1, max_c + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            st = cell_style(cell)
            cls = []
            if cell.alignment is not None and cell.alignment.wrap_text:
                cls.append("wrap")
            sticky = []
            if r <= frz_r:
                sticky.append(f"top:{top}px")
            if c <= frz_c:
                sticky.append(f"left:{lefts[c - 1]}px")
            if sticky:
                cls.append("frz")
                st = ";".join(filter(None, [st, *sticky, "background:" + (re.search(r"background:([^;]+)", st).group(1) if "background:" in st else "#fff")]))
            span = ""
            if (r, c) in merged:
                rs, cs = merged[(r, c)]
                span = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
                if cs > 1:
                    cls.append("wrap") if cell.alignment is not None and cell.alignment.wrap_text else None
            text = html.escape(show(cell)).replace("\n", "<br>")
            if cell.hyperlink is not None and cell.hyperlink.target:
                text = f'<a href="{html.escape(cell.hyperlink.target)}" target="_blank" rel="noopener">{text}</a>'
            attrs = (f' class="{" ".join(cls)}"' if cls else "") + (f' style="{st}"' if st else "") + span
            out.append(f"<td{attrs}>{text}</td>")
        out.append("</tr>")
        if r <= frz_r:
            top += h
    out.append("</tbody></table>")
    return "".join(out)


def xlsx_preview(path):
    wb = openpyxl.load_workbook(path)
    parts = [f"<style>{XL_CSS}</style>", '<div class="xl">']
    for i, ws in enumerate(wb.worksheets):
        if ws.sheet_state != "visible":
            continue
        parts.append(f'<section class="sheet{" on" if i == 0 else ""}" data-name="{html.escape(ws.title)}">{sheet_html(ws)}</section>')
    parts.append("</div>")
    return "".join(parts)


# ---------------------------------------------------------------- docx
PAGE_CSS = """
:host { all: initial; }
.paper-wrap { padding: 28px 0 40px; display: flex; justify-content: center; min-width: max-content; }
.paper { background: #fff; width: 8.5in; box-sizing: border-box; padding: 0.9in 1in; box-shadow: 0 1px 4px rgba(0,0,0,.18); color: #000; }
.paper img { max-width: 100%; height: auto; }
"""


def docx_preview(path):
    import pythoncom
    import win32com.client as win32
    pythoncom.CoInitialize()
    tmp = pathlib.Path(tempfile.mkdtemp())
    word = win32.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
        out = tmp / "doc.htm"
        doc.SaveAs2(str(out), FileFormat=10)  # wdFormatFilteredHTML
        doc.Close(False)
    finally:
        word.Quit()
    raw = out.read_bytes()
    m = re.search(rb'charset=([\w-]+)', raw)
    text = raw.decode(m.group(1).decode() if m else "windows-1252", errors="replace")
    styles = "\n".join(re.findall(r"<style[^>]*>(.*?)</style>", text, re.S | re.I))
    styles = re.sub(r"<!--|-->", "", styles)
    styles = re.sub(r"(?<![\w-])(body|html)\s*\{", ".paper {", styles)
    body = re.search(r"<body[^>]*>(.*)</body>", text, re.S | re.I).group(1)
    # images Word wrote beside the file become inline data
    import base64
    def inline(mm):
        src = html.unescape(mm.group(1))
        f = tmp / src.replace("/", "\\")
        if f.exists():
            ext = f.suffix.lstrip(".").lower().replace("jpg", "jpeg")
            return 'src="data:image/%s;base64,%s"' % (ext, base64.b64encode(f.read_bytes()).decode())
        return mm.group(0)
    body = re.sub(r'src="([^"]+)"', inline, body)
    shutil.rmtree(tmp, ignore_errors=True)
    return f"<style>{PAGE_CSS}\n{styles}</style><div class=\"paper-wrap\"><div class=\"paper\">{body}</div></div>"


# ---------------------------------------------------------------- html
def html_preview(path):
    text = path.read_text(encoding="utf-8")
    styles = "\n".join(re.findall(r"<style[^>]*>(.*?)</style>", text, re.S | re.I))
    styles = re.sub(r"(?<![\w-])body\s*\{", ".mail {", styles)
    body = re.search(r"<body[^>]*>(.*)</body>", text, re.S | re.I).group(1)
    return f"<style>:host {{ all: initial; }} .mail-wrap {{ background:#fff; min-width: max-content; }}\n{styles}</style><div class=\"mail-wrap\"><div class=\"mail\">{body}</div></div>"


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for old in OUT.glob("*.png"):
        old.unlink()
    for key, name in FILES.items():
        src = SAMPLES / name
        if src.suffix == ".xlsx":
            out = xlsx_preview(src)
        elif src.suffix == ".docx":
            out = docx_preview(src)
        else:
            out = html_preview(src)
        (OUT / f"{key}.html").write_text(out, encoding="utf-8")
        print(f"  {key:18s} {len(out) // 1024:4d} KB  from {name}")


if __name__ == "__main__":
    main()
